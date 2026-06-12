"""
从 resume_data.xlsx 导入前 N 条简历到学生表
用法:
    python import_resumes.py                    # 导入前 20 条（默认）
    python import_resumes.py --limit 50         # 导入前 50 条
    python import_resumes.py --limit 20 --skip-existing  # 跳过已存在的
    python import_resumes.py --limit 20 --dry-run         # 只预览不写入
"""
import asyncio
import argparse
import json
import sys
from pathlib import Path

# 确保能引用项目模块
sys.path.insert(0, str(Path(__file__).parent))

from db.database import async_session
from db.models import Student


def parse_skills(row: dict) -> dict:
    """将 Excel 技能列解析为结构化技能字典"""
    categories = {
        "编程语言": ("programming_languages", row.get("编程语言", ""), row.get("编程语言熟练度", "")),
        "前端技术": ("frontend", row.get("前端技术", ""), row.get("前端技术熟练度", "")),
        "后端技术": ("backend", row.get("后端技术", ""), row.get("后端技术熟练度", "")),
        "数据库": ("database", row.get("数据库", ""), row.get("数据库熟练度", "")),
        "云计算_运维": ("cloud_ops", row.get("云计算/运维", ""), row.get("云计算/运维熟练度", "")),
        "数据与算法": ("data_algo", row.get("数据与算法", ""), row.get("数据与算法熟练度", "")),
        "移动开发": ("mobile", row.get("移动开发", ""), row.get("移动开发熟练度", "")),
        "测试工具": ("testing", row.get("测试工具", ""), row.get("测试工具熟练度", "")),
    }

    result = {}
    for name, (key, techs, proficiency) in categories.items():
        if techs and str(techs).strip():
            tech_list = [t.strip() for t in str(techs).split(",") if t.strip()]
            result[key] = {
                "name": name,
                "items": tech_list,
                "proficiency": str(proficiency).strip() if proficiency else "",
            }
    return result


def parse_projects(row: dict) -> list:
    """解析工作经验 + 项目经历"""
    projects = []

    # 工作经验
    exp_fields = [
        ("小型企业工作经验", "小企业"),
        ("中型企业工作经验", "中企业"),
        ("大型企业工作经验", "大企业"),
    ]
    for field, scale in exp_fields:
        val = row.get(field, "")
        if val and str(val).strip() and str(val).strip() != "0":
            projects.append({
                "type": "work_experience",
                "scale": scale,
                "description": str(val).strip(),
            })

    # 项目经历
    proj_fields = [
        ("小规模项目", "小型"),
        ("中规模项目", "中型"),
        ("大规模项目", "大型"),
    ]
    for field, scale in proj_fields:
        val = row.get(field, "")
        if val and str(val).strip() and str(val).strip() != "0":
            projects.append({
                "type": "project",
                "scale": scale,
                "description": str(val).strip(),
            })

    return projects


def parse_academic(row: dict) -> dict:
    """解析学业基础"""
    return {
        "education_level": str(row.get("学历层次", "")).strip(),
        "school_type": str(row.get("院校类别", "")).strip(),
        "major_category": str(row.get("专业类别", "")).strip(),
        "english_level": str(row.get("英语水平", "")).strip(),
    }


def build_resume_text(row: dict, skills: dict, projects: list, academic: dict) -> str:
    """根据所有字段生成简历全文，供 AI 诊断使用"""
    parts = []

    name = str(row.get("姓名", "")).strip()
    gender = str(row.get("性别", "")).strip()
    age = str(row.get("年龄", "")).strip()
    target = str(row.get("意向岗位", "")).strip()

    parts.append(f"姓名：{name}，性别：{gender}，年龄：{age}")
    parts.append(f"意向岗位：{target}")
    parts.append(f"学历：{academic.get('education_level', '')}，院校类别：{academic.get('school_type', '')}，专业类别：{academic.get('major_category', '')}")
    parts.append(f"英语水平：{academic.get('english_level', '')}")

    # 技能
    skill_lines = []
    for cat_key, cat_data in skills.items():
        items = cat_data.get("items", [])
        prof = cat_data.get("proficiency", "")
        if items:
            skill_lines.append(f"{cat_data['name']}：{', '.join(items)}（{prof}）")
    if skill_lines:
        parts.append("技术技能：\n" + "\n".join(skill_lines))

    # 项目经验
    if projects:
        proj_lines = []
        for p in projects:
            proj_lines.append(f"[{p['type']}][{p['scale']}] {p['description']}")
        parts.append("项目/工作经验：\n" + "\n".join(proj_lines))

    return "\n\n".join(parts)


async def import_resumes(filepath: str, limit: int = 20, skip_existing: bool = False, dry_run: bool = False):
    """主导入逻辑"""
    import openpyxl

    wb = openpyxl.load_workbook(filepath)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]

    rows = []
    for row_cells in ws.iter_rows(min_row=2, max_row=min(limit + 1, ws.max_row), values_only=True):
        row_dict = dict(zip(headers, row_cells))
        # 跳过空行
        if not row_dict.get("姓名"):
            continue
        rows.append(row_dict)

    print(f"读取到 {len(rows)} 条简历（limit={limit}）")

    if dry_run:
        print("\n=== 预览模式（不写入数据库）===\n")
        for i, row in enumerate(rows):
            name = row.get("姓名", "")
            target = row.get("意向岗位", "")
            phone = row.get("电话", "")
            email = row.get("邮箱", "")
            edu = row.get("学历层次", "")
            print(f"[{i+1}] {name} | {target} | {edu} | {phone} | {email}")
        return

    imported = 0
    skipped = 0

    async with async_session() as session:
        from sqlalchemy import select

        for i, row in enumerate(rows):
            name = str(row.get("姓名", "")).strip()
            phone = str(row.get("电话", "")).strip()
            email = str(row.get("邮箱", "")).strip()

            # 跳过已存在的（按姓名+电话判断）
            if skip_existing:
                result = await session.execute(
                    select(Student).where(Student.name == name, Student.phone == phone)
                )
                existing = result.scalars().first()
                if existing:
                    skipped += 1
                    print(f"[{i+1}] 跳过已存在: {name}")
                    continue

            skills = parse_skills(row)
            projects = parse_projects(row)
            academic = parse_academic(row)
            resume_text = build_resume_text(row, skills, projects, academic)

            student = Student(
                name=name,
                phone=phone,
                email=email,
                target_job=str(row.get("意向岗位", "")).strip(),
                education_level=str(row.get("学历层次", "")).strip(),
                school=str(row.get("院校类别", "")).strip(),
                major=str(row.get("专业类别", "")).strip(),
                tech_skills=skills,
                project_exp=projects,
                academic_foundation=academic,
                resume_text=resume_text,
                soft_skills={},
                domain_knowledge={},
                soft_skill_evidence={},
                profile_sections={},
                profile_completeness=60.0,
            )

            session.add(student)
            imported += 1
            print(f"[{i+1}] 导入成功: {name} ({student.target_job})")

        await session.commit()

    print(f"\n导入完成：成功 {imported} 条，跳过 {skipped} 条")
    return imported


def main():
    parser = argparse.ArgumentParser(description="从 Excel 导入简历数据到学生表")
    parser.add_argument("--file", "-f", default=None, help="Excel 文件路径（默认：resume_data.xlsx）")
    parser.add_argument("--limit", "-n", type=int, default=20, help="导入前 N 条（默认 20）")
    parser.add_argument("--skip-existing", action="store_true", help="跳过已存在的同名+同电话学生")
    parser.add_argument("--dry-run", action="store_true", help="只预览不写入数据库")
    args = parser.parse_args()

    filepath = args.file or str(Path(__file__).resolve().parent / "resume_data.xlsx")
    if not Path(filepath).exists():
        print(f"❌ 文件不存在: {filepath}")
        sys.exit(1)

    asyncio.run(import_resumes(filepath, args.limit, args.skip_existing, args.dry_run))


if __name__ == "__main__":
    main()
